"""
tests/test_ppt_evaluation_service.py

Phase 5 - Tests for ppt_evaluation_service.py

Test strategy
-------------
Unit tests (1-14): Gemini is MOCKED — deterministic, no API key required.
Integration test (15): Real Gemini call using NIRMAAN_template.pdf + synthetic
student PPT from Phase 4. Skipped on missing API key or Gemini 503.
Regression test (16): Verifies Project Validation pipeline unchanged.

Tests
-----
 1. Import and syntax check
 2. Strong complete presentation -> high scores
 3. Weak presentation -> low scores
 4. Missing template (Mode B) -> Template Compliance = 0, weight redistributed
 5. Template compliance score reflects RequirementMatch results
 6. Missing requirements generate findings
 7. Partial requirements reduce Template Compliance
 8. Slide analysis generated for every slide
 9. Evidence is preserved in criterion scores
10. No fabricated evidence (NOT FOUND explicitly stated)
11. Invalid Gemini JSON -> ValueError
12. Overall score bounds 0-100
13. Criterion score bounds 0-10
14. Judge questions generated from gaps
15. Real Gemini integration: NIRMAAN + synthetic PPT
16. Project Validation pipeline unaffected
"""

from __future__ import annotations

import asyncio
import io
import os
import sys
import traceback
import unittest.mock as mock
from typing import List

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.schemas.evaluation import FindingType, Severity
from app.schemas.ppt import (
    MatchStatus,
    PPTCriterionScore,
    PPTEvaluationResponse,
    PPTFinding,
    RequirementMatch,
    RequirementType,
    SlideAnalysis,
    TemplateAnalysis,
    TemplateRequirement,
)
from app.services.extraction_service import (
    ExtractionResult,
    FileType,
    PresentationMetadata,
    SlideExtraction,
    TextBlock,
)
from app.services.ppt_evaluation_service import (
    CRITERIA,
    WEIGHTS_WITH_TEMPLATE,
    WEIGHTS_WITHOUT_TEMPLATE,
    _GeminiEvalResponse,
    _SlideAnalysisItem,
    _CriterionScoreItem,
    _FindingItem,
    _compute_overall_score,
    _template_compliance_from_matches,
    evaluate_presentation,
)

# ---------------------------------------------------------------------------
# Runner helpers
# ---------------------------------------------------------------------------

PASS = "\033[92mPASS\033[0m"
FAIL = "\033[91mFAIL\033[0m"
SKIP = "\033[93mSKIP\033[0m"
results: list[tuple[str, str, str | None]] = []


class SkipTest(Exception):
    pass


def run_test(name: str, fn):
    try:
        fn()
        print(f"  [{PASS}] {name}")
        results.append((name, "pass", None))
    except SkipTest as e:
        print(f"  [{SKIP}] {name}: {e}")
        results.append((name, "skip", str(e)))
    except AssertionError as e:
        print(f"  [{FAIL}] {name}: {e}")
        results.append((name, "fail", str(e)))
    except Exception as e:
        print(f"  [{FAIL}] {name}: {type(e).__name__}: {e}")
        print(f"         {traceback.format_exc()}")
        results.append((name, "fail", f"{type(e).__name__}: {e}"))


def run_async_test(name: str, coro_fn):
    try:
        asyncio.run(coro_fn())
        print(f"  [{PASS}] {name}")
        results.append((name, "pass", None))
    except SkipTest as e:
        print(f"  [{SKIP}] {name}: {e}")
        results.append((name, "skip", str(e)))
    except AssertionError as e:
        print(f"  [{FAIL}] {name}: {e}")
        results.append((name, "fail", str(e)))
    except Exception as e:
        print(f"  [{FAIL}] {name}: {type(e).__name__}: {e}")
        print(f"         {traceback.format_exc()}")
        results.append((name, "fail", f"{type(e).__name__}: {e}"))


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _meta(n: int = 3) -> PresentationMetadata:
    return PresentationMetadata(file_type=FileType.pdf, total_slides=n)


def _slide(number: int, title: str = "", text: str = "", images: int = 0) -> SlideExtraction:
    blocks = [TextBlock(shape_type="BODY", text=text)] if text else []
    return SlideExtraction(
        slide_number=number, title=title, text=text,
        text_blocks=blocks, image_count=images,
        shape_count=len(blocks) + (1 if title else 0),
        text_character_count=len(text),
    )


def _extraction(*slides) -> ExtractionResult:
    n = max((s.slide_number for s in slides), default=0)
    return ExtractionResult(metadata=_meta(n), slides=list(slides))


def _req(rid: str, title: str) -> TemplateRequirement:
    return TemplateRequirement(
        requirement_id=rid, title=title,
        description=f"Requirement {rid}",
        requirement_type=RequirementType.required,
        keywords=[], source="Test template",
    )


def _match(rid: str, title: str, status: MatchStatus, slides=None) -> RequirementMatch:
    return RequirementMatch(
        requirement_id=rid, requirement_title=title,
        status=status,
        matching_slides=slides or [],
        explanation=f"Test explanation for {rid}",
        recommendations=[] if status == MatchStatus.matched else [f"Improve {title}"],
    )


def _template(*reqs) -> TemplateAnalysis:
    return TemplateAnalysis(
        template_name="Test Template", total_slides=7,
        requirements=list(reqs),
        overall_structure=[], notes=[],
    )


def _all_criterion_scores(
    scores: dict[str, float] | None = None,
    default: float = 7.0,
) -> List[_CriterionScoreItem]:
    scores = scores or {}
    return [
        _CriterionScoreItem(
            criterion=c,
            score=scores.get(c, default),
            explanation=f"Test explanation for {c}",
            evidence=[f"Evidence for {c}"],
        )
        for c in CRITERIA
    ]


def _all_slides_analysis(extraction: ExtractionResult) -> List[_SlideAnalysisItem]:
    return [
        _SlideAnalysisItem(
            slide_number=s.slide_number,
            title=s.title or "",
            summary=f"Summary of slide {s.slide_number}",
            content_present=bool(s.text),
            visual_elements_detected=[f"{s.image_count} images"] if s.image_count else [],
            strengths=["Good content"],
            weaknesses=["Could be improved"],
        )
        for s in extraction.slides
    ]


def _mock_gemini(
    extraction: ExtractionResult,
    scores: dict[str, float] | None = None,
    findings: list | None = None,
    strengths: list | None = None,
    weaknesses: list | None = None,
    judge_questions: list | None = None,
) -> mock.MagicMock:
    raw = _GeminiEvalResponse(
        overall_summary="Test evaluation summary.",
        scores=_all_criterion_scores(scores),
        slide_analysis=_all_slides_analysis(extraction),
        strengths=strengths or ["Strong problem definition.", "Clear solution."],
        weaknesses=weaknesses or ["Validation missing.", "Architecture detail lacking."],
        recommendations=["Add validation data.", "Clarify architecture."],
        judge_questions=judge_questions or [
            "How was the model validated?",
            "What is the deployment strategy?",
        ],
        findings=findings or [
            _FindingItem(
                finding_type=FindingType.missing,
                title="Evidence / Validation absent",
                description="No quantified validation found.",
                severity=Severity.high,
                source="Entire presentation",
            )
        ],
    )
    m = mock.MagicMock()
    m.parsed = raw
    m.text = None
    return m


# ---------------------------------------------------------------------------
# Test 1: Import check
# ---------------------------------------------------------------------------

def test_import():
    from app.services.ppt_evaluation_service import evaluate_presentation
    assert callable(evaluate_presentation)


# ---------------------------------------------------------------------------
# Test 2: Strong complete presentation -> high overall score
# ---------------------------------------------------------------------------

async def test_strong_presentation():
    reqs = [_req("R1", "Problem"), _req("R2", "Solution"), _req("R3", "Architecture")]
    tmpl = _template(*reqs)
    matches = [
        _match("R1", "Problem", MatchStatus.matched, [1]),
        _match("R2", "Solution", MatchStatus.matched, [2]),
        _match("R3", "Architecture", MatchStatus.matched, [3]),
    ]
    ext = _extraction(
        _slide(1, "Problem", "Clear real-world problem with target users identified."),
        _slide(2, "Solution", "Innovative AI-powered solution with proven approach."),
        _slide(3, "Architecture", "React, FastAPI, PostgreSQL, Docker deployment.", images=1),
    )
    # All matched -> compliance = 10/10
    expected_compliance = _template_compliance_from_matches(matches)
    assert expected_compliance == 10.0

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, scores={c: 9.0 for c in CRITERIA}
        )
        result = await evaluate_presentation(tmpl, matches, ext)

    assert isinstance(result, PPTEvaluationResponse)
    assert result.overall_score > 70.0, f"Strong PPT score too low: {result.overall_score}"
    assert result.template_compliance_score == 100.0
    assert len(result.scores) == 10
    print(f"       Strong PPT: overall={result.overall_score:.1f}, "
          f"compliance={result.template_compliance_score:.0f}/100")


# ---------------------------------------------------------------------------
# Test 3: Weak presentation -> low overall score
# ---------------------------------------------------------------------------

async def test_weak_presentation():
    reqs = [_req("R1", "Problem"), _req("R2", "Solution")]
    tmpl = _template(*reqs)
    matches = [
        _match("R1", "Problem", MatchStatus.missing),
        _match("R2", "Solution", MatchStatus.missing),
    ]
    ext = _extraction(_slide(1, "Intro", "We are building something."))

    expected_compliance = _template_compliance_from_matches(matches)
    assert expected_compliance == 0.0

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, scores={c: 2.0 for c in CRITERIA}
        )
        result = await evaluate_presentation(tmpl, matches, ext)

    assert result.overall_score < 40.0, f"Weak PPT score too high: {result.overall_score}"
    assert result.template_compliance_score == 0.0
    print(f"       Weak PPT: overall={result.overall_score:.1f}, "
          f"compliance={result.template_compliance_score:.0f}/100")


# ---------------------------------------------------------------------------
# Test 4: No template (Mode B) -> Template Compliance = 0, weight redistributed
# ---------------------------------------------------------------------------

async def test_no_template_mode():
    ext = _extraction(
        _slide(1, "Problem", "Water quality monitoring problem."),
        _slide(2, "Solution", "IoT sensors + AI."),
    )

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, scores={c: 7.0 for c in CRITERIA}
        )
        result = await evaluate_presentation(None, [], ext)

    # Template Compliance must be 0 in no-template mode
    assert result.template_compliance_score == 0.0
    tc_score = next((s for s in result.scores if s.criterion == "Template Compliance"), None)
    assert tc_score is not None
    assert tc_score.score == 0.0, f"Template Compliance should be 0, got {tc_score.score}"

    # Weights without template sum to 100% across 9 criteria
    assert abs(sum(WEIGHTS_WITHOUT_TEMPLATE.values()) - 100.0) < 0.01, \
        "WEIGHTS_WITHOUT_TEMPLATE must sum to 100%"

    # Overall score should still be computed (not zero)
    assert result.overall_score > 0.0, "Overall score should be > 0 even without template"

    print(f"       No-template: overall={result.overall_score:.1f}, "
          f"compliance=0/100 (correct)")


# ---------------------------------------------------------------------------
# Test 5: Template compliance score reflects RequirementMatch results
# ---------------------------------------------------------------------------

def test_compliance_score_from_matches():
    # 3 matched, 1 partially, 2 missing = (3*1 + 1*0.5 + 2*0) / 6 = 3.5/6 * 10 = 5.83
    matches = [
        _match("R1", "A", MatchStatus.matched),
        _match("R2", "B", MatchStatus.matched),
        _match("R3", "C", MatchStatus.matched),
        _match("R4", "D", MatchStatus.partially_matched),
        _match("R5", "E", MatchStatus.missing),
        _match("R6", "F", MatchStatus.missing),
    ]
    score = _template_compliance_from_matches(matches)
    expected = round((3.5 / 6) * 10, 2)
    assert abs(score - expected) < 0.01, f"Expected {expected}, got {score}"

    # All not_applicable -> 10.0
    na_matches = [_match("R1", "A", MatchStatus.not_applicable)]
    assert _template_compliance_from_matches(na_matches) == 10.0

    # Empty -> 0.0
    assert _template_compliance_from_matches([]) == 0.0

    print(f"       Compliance calculation: 3 matched + 1 partial + 2 missing = {score:.2f}/10 [OK]")


# ---------------------------------------------------------------------------
# Test 6: Missing requirements generate findings
# ---------------------------------------------------------------------------

async def test_missing_requirements_in_findings():
    reqs = [_req("R1", "Demo Plan"), _req("R2", "Architecture")]
    tmpl = _template(*reqs)
    matches = [
        _match("R1", "Demo Plan", MatchStatus.missing),
        _match("R2", "Architecture", MatchStatus.matched, [2]),
    ]
    ext = _extraction(_slide(1, "Intro", "Text"), _slide(2, "Architecture", "Tech stack listed"))

    missing_finding = _FindingItem(
        finding_type=FindingType.missing,
        title="Demo Plan absent",
        description="No demo plan slide found.",
        severity=Severity.high,
        source="Entire presentation",
    )

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, findings=[missing_finding]
        )
        result = await evaluate_presentation(tmpl, matches, ext)

    assert len(result.findings) > 0
    finding_titles = [f.title for f in result.findings]
    assert any("Demo" in t or "demo" in t for t in finding_titles), \
        f"Expected Demo Plan finding, got: {finding_titles}"
    print(f"       Findings generated: {finding_titles}")


# ---------------------------------------------------------------------------
# Test 7: Partial requirements reduce Template Compliance
# ---------------------------------------------------------------------------

def test_partial_reduces_compliance():
    all_partial = [_match(f"R{i}", f"Req{i}", MatchStatus.partially_matched) for i in range(5)]
    score = _template_compliance_from_matches(all_partial)
    assert score == 5.0, f"All partially_matched should give 5.0, got {score}"

    all_matched = [_match(f"R{i}", f"Req{i}", MatchStatus.matched) for i in range(5)]
    assert _template_compliance_from_matches(all_matched) == 10.0

    print(f"       All partial -> 5.0/10 [OK], All matched -> 10.0/10 [OK]")


# ---------------------------------------------------------------------------
# Test 8: Slide analysis generated for every slide
# ---------------------------------------------------------------------------

async def test_slide_analysis_coverage():
    ext = _extraction(
        _slide(1, "Slide One", "Text A"),
        _slide(2, "Slide Two", "Text B"),
        _slide(3, "Slide Three", "Text C"),
    )

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        # Gemini returns analysis only for slides 1 and 2 (slide 3 missing)
        gemini_resp = _GeminiEvalResponse(
            overall_summary="Summary",
            scores=_all_criterion_scores(),
            slide_analysis=[
                _SlideAnalysisItem(
                    slide_number=1, title="Slide One", summary="S1",
                    content_present=True,
                    visual_elements_detected=[], strengths=[], weaknesses=[],
                ),
                _SlideAnalysisItem(
                    slide_number=2, title="Slide Two", summary="S2",
                    content_present=True,
                    visual_elements_detected=[], strengths=[], weaknesses=[],
                ),
                # slide 3 intentionally omitted — fallback entry should be added
            ],
            strengths=["Good"], weaknesses=["Needs work"],
            recommendations=["Improve"], judge_questions=["Why?"], findings=[],
        )
        m = mock.MagicMock()
        m.parsed = gemini_resp
        gc.return_value.models.generate_content.return_value = m

        result = await evaluate_presentation(None, [], ext)

    slide_numbers = [sa.slide_number for sa in result.slide_analysis]
    assert 1 in slide_numbers
    assert 2 in slide_numbers
    assert 3 in slide_numbers, "Slide 3 should have a fallback entry"
    assert len(result.slide_analysis) == 3
    print(f"       Slide analysis covers all 3 slides including fallback for slide 3 [OK]")


# ---------------------------------------------------------------------------
# Test 9: Evidence preserved in criterion scores
# ---------------------------------------------------------------------------

async def test_evidence_preserved():
    ext = _extraction(_slide(1, "Tech Stack", "React, FastAPI, PostgreSQL"))

    evidence_items = ["React frontend detected", "FastAPI backend mentioned", "PostgreSQL database listed"]

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        resp = _GeminiEvalResponse(
            overall_summary="Summary",
            scores=[
                _CriterionScoreItem(
                    criterion=c,
                    score=7.0,
                    explanation=f"Explanation for {c}",
                    evidence=evidence_items if c == "Technical Depth" else [],
                )
                for c in CRITERIA
            ],
            slide_analysis=_all_slides_analysis(ext),
            strengths=["Good tech stack"],
            weaknesses=[], recommendations=[], judge_questions=[], findings=[],
        )
        m = mock.MagicMock()
        m.parsed = resp
        gc.return_value.models.generate_content.return_value = m

        result = await evaluate_presentation(None, [], ext)

    tech_score = next(s for s in result.scores if s.criterion == "Technical Depth")
    assert tech_score.evidence == evidence_items, \
        f"Evidence not preserved: {tech_score.evidence}"
    print(f"       Evidence preserved in Technical Depth: {evidence_items}")


# ---------------------------------------------------------------------------
# Test 10: NOT FOUND must be stated when evidence is absent
# ---------------------------------------------------------------------------

async def test_not_found_for_missing_evidence():
    ext = _extraction(_slide(1, "Introduction", "We are building a water monitor."))

    validation_explanation = "NOT FOUND: No user testing, benchmark, pilot, or validation data found."
    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        scores_with_not_found = _all_criterion_scores()
        for s in scores_with_not_found:
            if s.criterion == "Evidence / Validation":
                s.score = 1.5
                s.explanation = validation_explanation
        resp = _GeminiEvalResponse(
            overall_summary="Summary",
            scores=scores_with_not_found,
            slide_analysis=_all_slides_analysis(ext),
            strengths=[], weaknesses=["No validation"],
            recommendations=["Add validation"], judge_questions=["Any data?"], findings=[],
        )
        m = mock.MagicMock()
        m.parsed = resp
        gc.return_value.models.generate_content.return_value = m

        result = await evaluate_presentation(None, [], ext)

    ev_score = next(s for s in result.scores if s.criterion == "Evidence / Validation")
    assert "NOT FOUND" in ev_score.explanation or ev_score.score <= 3.0, \
        f"Validation explanation should indicate NOT FOUND: {ev_score.explanation}"
    print(f"       Evidence/Validation score={ev_score.score:.1f}: {ev_score.explanation[:60]}...")


# ---------------------------------------------------------------------------
# Test 11: Invalid Gemini JSON -> ValueError
# ---------------------------------------------------------------------------

async def test_invalid_gemini_json():
    ext = _extraction(_slide(1, "Slide", "Content"))

    bad = mock.MagicMock()
    bad.parsed = None
    bad.text = '{"this": "is not valid _GeminiEvalResponse"}'

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = bad
        try:
            await evaluate_presentation(None, [], ext)
            assert False, "Should have raised ValueError"
        except ValueError as e:
            assert "invalid JSON" in str(e).lower() or "validation" in str(e).lower(), \
                f"Wrong error message: {e}"
    print(f"       Invalid Gemini JSON correctly raises ValueError [OK]")


# ---------------------------------------------------------------------------
# Test 12: Overall score always in 0-100
# ---------------------------------------------------------------------------

def test_score_bounds_overall():
    # Test _compute_overall_score with extreme inputs
    scores_max = [PPTCriterionScore(criterion=c, score=10.0, explanation="max") for c in CRITERIA]
    scores_min = [PPTCriterionScore(criterion=c, score=0.0, explanation="min") for c in CRITERIA]
    scores_mid = [PPTCriterionScore(criterion=c, score=5.0, explanation="mid") for c in CRITERIA]

    assert _compute_overall_score(scores_max, True) == 100.0
    assert _compute_overall_score(scores_min, True) == 0.0
    mid = _compute_overall_score(scores_mid, True)
    assert 49.0 <= mid <= 51.0, f"Mid scores should give ~50, got {mid}"

    # Without template: Template Compliance excluded
    scores_no_tmpl = [
        PPTCriterionScore(criterion=c, score=10.0, explanation="max")
        for c in CRITERIA if c != "Template Compliance"
    ]
    assert _compute_overall_score(scores_no_tmpl, False) == 100.0

    print(f"       Score bounds: max=100.0, min=0.0, mid={mid:.1f} [OK]")


# ---------------------------------------------------------------------------
# Test 13: Criterion scores always 0-10
# ---------------------------------------------------------------------------

async def test_criterion_score_bounds():
    """
    Verify that _enforce_score_bounds clamps out-of-range scores.
    Uses model_construct() to bypass Pydantic validation so we can
    simulate what would happen if Gemini returned an out-of-bounds score
    (which the real API prevents, but we test our post-processing guard).
    """
    from app.services.ppt_evaluation_service import _enforce_score_bounds

    ext = _extraction(_slide(1, "Slide", "Content"))

    # Use model_construct to bypass Pydantic ge/le validation intentionally
    out_of_bounds_scores = [
        _CriterionScoreItem.model_construct(
            criterion=c, score=12.0,
            explanation="out of bounds", evidence=[],
        )
        for c in CRITERIA
    ]

    # Directly test the post-processing function
    clamped = _enforce_score_bounds(out_of_bounds_scores)
    for s in clamped:
        assert s.score <= 10.0, f"Score not clamped: {s.criterion}={s.score}"
        assert s.score >= 0.0, f"Score below zero: {s.criterion}={s.score}"

    # Also verify via the full pipeline using a mock that returns valid scores
    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, scores={c: 7.0 for c in CRITERIA}
        )
        result = await evaluate_presentation(None, [], ext)

    for s in result.scores:
        assert 0.0 <= s.score <= 10.0, f"Score out of bounds: {s.criterion}={s.score}"

    print(f"       _enforce_score_bounds clamps 12.0 -> 10.0 [OK], pipeline scores verified [OK]")


# ---------------------------------------------------------------------------
# Test 14: Judge questions generated from gaps
# ---------------------------------------------------------------------------

async def test_judge_questions_generated():
    ext = _extraction(_slide(1, "Problem", "Vague problem statement."))
    questions = [
        "How was the solution validated?",
        "What is the deployment strategy?",
        "How does this scale beyond a prototype?",
    ]

    with mock.patch("app.services.ppt_evaluation_service.get_client") as gc:
        gc.return_value.models.generate_content.return_value = _mock_gemini(
            ext, judge_questions=questions
        )
        result = await evaluate_presentation(None, [], ext)

    assert len(result.judge_questions) >= 1
    assert result.judge_questions == questions
    print(f"       Judge questions: {result.judge_questions}")


# ---------------------------------------------------------------------------
# Test 15: Real Gemini integration — NIRMAAN + synthetic PPT
# ---------------------------------------------------------------------------

def _make_synthetic_pptx_bytes() -> bytes:
    """Same 5-slide synthetic PPT as Phase 4 (with intentional gaps)."""
    from pptx import Presentation
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slides_data = [
        {"title": "AquaSense — Smart Water Quality Monitor",
         "body": "Team Name: AquaSense  |  Team Leader: Riya Sharma\n"
                 "Chosen Track: Environmental Tech  |  Hackathon: NIRMAAN 2026"},
        {"title": "Problem Statement & Real-World Impact",
         "body": "2B+ people lack safe drinking water. Manual monitoring is slow.\n"
                 "Target: Rural communities, water boards, NGOs."},
        {"title": "Proposed Solution & Innovation",
         "body": "IoT sensor array + edge AI + cloud dashboard.\n"
                 "Innovation: On-device inference reduces latency 80%."},
        {"title": "Architecture",   # Intentionally weak
         "body": "We use sensors connected to the cloud."},
        {"title": "Team Details",
         "body": "Riya Sharma (21CS101), Aarav Patel (21EC202)\n"
                 "Open-source libraries: MIT/Apache licensed."},
    ]
    for sd in slides_data:
        s = prs.slides.add_slide(layout)
        if s.shapes.title:
            s.shapes.title.text = sd["title"]
        try:
            s.placeholders[1].text = sd["body"]
        except Exception:
            pass
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_real_gemini_integration():
    """Full pipeline: NIRMAAN template -> template analysis -> requirement matching
    -> PPT evaluation. Skipped on missing API key or Gemini 503."""
    from app.services.gemini_service import get_client as _gc
    try:
        _gc()
    except RuntimeError:
        raise SkipTest("GEMINI_API_KEY not configured")

    from app.services.extraction_service import extract_document, extract_pptx
    from app.services.template_analyzer import analyze_template
    from app.services.requirement_matcher import match_requirements

    nirmaan = os.path.join(os.path.dirname(__file__), "NIRMAAN_template.pdf")
    if not os.path.exists(nirmaan):
        raise SkipTest("NIRMAAN_template.pdf not found")

    with open(nirmaan, "rb") as f:
        template_bytes = f.read()

    print("\n       [1] Extracting NIRMAAN template...")
    tmpl_extraction = extract_document(template_bytes, "NIRMAAN_template.pdf")

    print("       [2] Analyzing template with Gemini...")
    try:
        tmpl_analysis = await analyze_template(tmpl_extraction, "NIRMAAN_template.pdf")
    except Exception as e:
        if "503" in str(e) or "UNAVAILABLE" in str(e) or "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e):
            raise SkipTest(f"Gemini 503 — try again later: {e}")
        raise

    print(f"            Requirements: {len(tmpl_analysis.requirements)}")

    print("       [3] Creating synthetic student PPT...")
    student_bytes = _make_synthetic_pptx_bytes()
    student_extraction = extract_pptx(student_bytes)
    print(f"            Student PPT: {student_extraction.metadata.total_slides} slides")

    print("       [4] Running requirement matching...")
    try:
        matches = await match_requirements(tmpl_analysis, student_extraction)
    except Exception as e:
        if "503" in str(e) or "UNAVAILABLE" in str(e) or "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e):
            raise SkipTest(f"Gemini 503 during matching: {e}")
        raise

    print(f"            Matches: {[(m.requirement_id, m.status.value) for m in matches]}")

    print("       [5] Running PPT evaluation...")
    try:
        result = await evaluate_presentation(tmpl_analysis, matches, student_extraction)
    except Exception as e:
        if "503" in str(e) or "UNAVAILABLE" in str(e) or "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e):
            raise SkipTest(f"Gemini 503 during evaluation: {e}")
        raise

    # --- Assertions ---
    assert isinstance(result, PPTEvaluationResponse)
    assert 0.0 <= result.overall_score <= 100.0
    assert 0.0 <= result.template_compliance_score <= 100.0
    assert len(result.scores) == 10
    assert len(result.slide_analysis) == student_extraction.metadata.total_slides

    for s in result.scores:
        assert 0.0 <= s.score <= 10.0, f"{s.criterion} score out of bounds: {s.score}"

    # Verify NIRMAAN-specific expectations (based on intentional gaps)
    missing_matches = [m for m in matches if m.status == MatchStatus.missing]
    assert len(missing_matches) >= 1, "Expected at least 1 missing requirement (REQ-005, REQ-006)"

    # Template compliance should NOT be 100 (has partially matched and missing)
    assert result.template_compliance_score < 100.0

    # Should have weaknesses/recommendations
    assert len(result.weaknesses) > 0
    assert len(result.recommendations) > 0
    assert len(result.judge_questions) > 0

    print(f"\n       RESULTS:")
    print(f"         Overall score:          {result.overall_score:.1f}/100")
    print(f"         Template compliance:    {result.template_compliance_score:.0f}/100")
    print(f"         Criteria scores:")
    for s in result.scores:
        print(f"           {s.criterion:35s} {s.score:.1f}/10")
    print(f"         Slide analysis:         {len(result.slide_analysis)} slides")
    print(f"         Strengths:              {len(result.strengths)}")
    print(f"         Weaknesses:             {len(result.weaknesses)}")
    print(f"         Recommendations:        {len(result.recommendations)}")
    print(f"         Judge questions:        {len(result.judge_questions)}")
    print(f"         Findings:               {len(result.findings)}")
    if result.findings:
        for f in result.findings[:3]:
            print(f"           [{f.severity.value:8s}/{f.finding_type.value:14s}] {f.title}")


# ---------------------------------------------------------------------------
# Test 16: Project Validation still works
# ---------------------------------------------------------------------------

def test_project_validation_unaffected():
    import app.main
    from app.schemas.evaluation import EvaluationResponse
    from app.services.evaluation_service import evaluate_project

    from app.main import app as fa
    routes = {r.path for r in fa.routes}
    assert "/api/projects/evaluate" in routes, f"Route missing: {routes}"
    print(f"       /api/projects/evaluate confirmed present [OK]")


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

def main():
    print("\n" + "=" * 68)
    print("  BuddyJudge Phase 5 - PPT Evaluation Service Tests")
    print("=" * 68 + "\n")

    run_test("Import check", test_import)
    run_async_test("Strong complete presentation -> high score", test_strong_presentation)
    run_async_test("Weak presentation -> low score", test_weak_presentation)
    run_async_test("No template (Mode B) -> compliance=0, weights redistributed", test_no_template_mode)
    run_test("Compliance score computed from RequirementMatch", test_compliance_score_from_matches)
    run_async_test("Missing requirements appear in findings", test_missing_requirements_in_findings)
    run_test("Partial matches reduce compliance to 5.0/10", test_partial_reduces_compliance)
    run_async_test("Slide analysis covers every slide (fallback)", test_slide_analysis_coverage)
    run_async_test("Evidence preserved in criterion scores", test_evidence_preserved)
    run_async_test("NOT FOUND stated for absent evidence", test_not_found_for_missing_evidence)
    run_async_test("Invalid Gemini JSON -> ValueError", test_invalid_gemini_json)
    run_test("Overall score bounds 0-100", test_score_bounds_overall)
    run_async_test("Criterion scores clamped to 0-10", test_criterion_score_bounds)
    run_async_test("Judge questions generated", test_judge_questions_generated)
    run_async_test("Real Gemini: NIRMAAN + synthetic PPT (full pipeline)", test_real_gemini_integration)
    run_test("Project Validation unaffected", test_project_validation_unaffected)

    print("\n" + "=" * 68)
    passed = sum(1 for _, s, _ in results if s == "pass")
    skipped = sum(1 for _, s, _ in results if s == "skip")
    failed = sum(1 for _, s, _ in results if s == "fail")
    total = len(results)
    print(f"  Results: {passed} passed / {skipped} skipped / {failed} failed  (total {total})")
    if failed == 0:
        print("  ALL TESTS PASSED (or skipped)")
    else:
        print("  FAILURES:")
        for name, status, err in results:
            if status == "fail":
                print(f"    - {name}: {err}")
    print("=" * 68 + "\n")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
